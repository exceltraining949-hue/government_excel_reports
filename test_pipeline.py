"""Direct pipeline test (no server)."""
import json, os, sys
from analyzer import load_file, analyze, to_payload
from report_builder import build_report

src = "/home/user/sample_messy_gov_data.xlsx"
sheets, notes, wb_meta = load_file(src, "sample_messy_gov_data.xlsx")
print("SHEETS:", {k: v.shape for k, v in sheets.items()})

ctx = {"file_name": "sample_messy_gov_data.xlsx", "sheet_name": "HR_Budget_Data",
       "auto_selected": True, "sheet_notes": notes, "hidden_sheets": [], "merged_ranges": {}}
result = analyze(sheets["HR_Budget_Data"], options={"standardize_case": True, "fill_missing": True, "remove_duplicates": False}, context=ctx)

print("\n--- KPIs ---")
for k, v in result["kpis"].items():
    print(f"  {k}: {v}")

print("\n--- ISSUES (%d) ---" % len(result["issues"]))
for i in result["issues"][:15]:
    print(f"  [{i['severity']}] {i['area']}: {i['description'][:100]}")

print("\n--- CHANGE LOG (%d) ---" % len(result["change_log"]))
for c in result["change_log"][:12]:
    print(f"  {c['field']}: {c['original'][:30]!r} -> {c['updated'][:30]!r} ({c['cells']} cells)")

print("\n--- PIVOTS ---")
for p in result["pivots"]:
    print(f"  {p['dimension']} (measure={p['measure']}): {len(p['rows'])} rows, top = {p['rows'][0]}")

print("\n--- BUDGET ---")
b = result["budget"]
if b:
    print(f"  budget={b['total_budget']:,.0f} expend={b['total_expenditure']:,.0f} util={b['utilization']}% above={b['above_budget_count']}")
    for d in b["per_dept"][:3]:
        print("   ", d)

print("\n--- TREND ---")
t = result["trend"]
if t:
    print("  months:", len(t["labels"]), t["labels"][:3], "...", t["labels"][-1])

print("\n--- OUTLIERS: total =", result["outlier_total"], ", listed =", len(result["outliers"]))
print("\n--- ASSUMPTIONS (%d) ---" % len(result["assumptions"]))
for a in result["assumptions"]:
    print("  -", a[:110])

print("\n--- TOP5 ---")
for r in result["top50"]["rows"][:5]:
    print("  ", r)

# JSON payload serialization check
payload = to_payload(result, "testjob", "sample_messy_gov_data.xlsx", "HR_Budget_Data",
                     [{"name": k, "rows": v.shape[0], "cols": v.shape[1], "hidden": False, "merged_ranges": 0} for k, v in sheets.items()],
                     "/download/testjob")
s = json.dumps(payload, ensure_ascii=False)
print("\nJSON payload OK:", len(s), "chars")

# Excel report build
out = "/home/user/test_report.xlsx"
build_report(result, out)
print("Report built:", os.path.getsize(out), "bytes")

# verify report opens
from openpyxl import load_workbook
wb = load_workbook(out)
print("Report sheets:", wb.sheetnames)
