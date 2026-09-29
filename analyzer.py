"""
GovData Analytics - Core Analysis Engine
=========================================
Implements the Government Office Excel Analytics workflow:
inspection -> profiling -> cleaning -> validation -> summaries ->
rankings -> trends -> anomaly detection -> management summary.

RULES ENFORCED HERE:
- Never fabricate data.
- Never silently modify records: every change is logged.
- Raw data is preserved; cleaning happens on a copy.
- Anomalies are flagged neutrally ("requires review").
"""

import re
import numpy as np
import pandas as pd

ERROR_STRINGS = {'#n/a', '#value!', '#ref!', '#div/0!', '#name?', '#num!', '#null!', 'n/a', 'na'}

# ------------------------------------------------------------------ #
# Column-role detection (name-pattern heuristics + dtype)
# ------------------------------------------------------------------ #
ROLE_ORDER = ['id', 'percent', 'budget', 'expenditure', 'date', 'department', 'district',
              'status', 'vendor', 'designation', 'grade', 'category', 'gender',
              'amount', 'name']

ROLE_PATTERNS = {
    # ID / code detection FIRST - a code must never be treated as a measure
    'id':           [r'\bid\b', r'\bno\.?$', r'_no\b', r'\bnumber\b', r'\bcode\b', r'\bref\b', r'\bkey\b',
                     r'application', r'registration', r'emp(loyee)?[_\s]*id', r'project[_\s]*id',
                     r'transaction[_\s]*id', r'invoice[_\s]*(no|id|#)', r'\bcmis\b', r'case[_\s]*(no|id|#)',
                     r'\bserial\b', r'\bsr\.?\s*#?$', r'\bcnic\b', r'\bpin\b'],
    'budget':       [r'budget', r'allocat', r'sanction', r'grant', r'approved\s*(cost|amount|fund)', r'\bbre\b'],
    'expenditure':  [r'expenditure', r'expense', r'spend', r'payment', r'\bpaid\b', r'utili[sz]', r'incurred', r'\bexpend\b'],
    'date':         [r'date', r'\bdob\b', r'join(ing|ed)?', r'retirement', r'completion', r'deadline', r'\bdue\b', r'tarikh', r'issued'],
    'department':   [r'department', r'\bdept\b', r'\bwing\b', r'branch', r'directorate', r'ministry', r'\bsection\b', r'\boffice\b'],
    'district':     [r'district', r'tehsil', r'zilla', r'\bcity\b', r'region', r'province', r'division', r'location', r'\barea\b'],
    'status':       [r'status', r'\bstage\b', r'\bstate\b', r'\bphase\b', r'progress', r'remark', r'\bcondition\b'],
    'vendor':       [r'vendor', r'supplier', r'contractor', r'\bfirm\b', r'company', r'bidder'],
    'designation':  [r'designation', r'\bpost\b', r'cadre', r'\bposition\b'],
    'grade':        [r'grade', r'scale', r'\bbps\b', r'\bpay\s*group\b'],
    'category':     [r'category', r'\btype\b', r'\bclass\b', r'\bgroup\b', r'scheme', r'sector'],
    'gender':       [r'gender', r'\bsex\b'],
    'percent':      [r'percent', r'%', r'\brate\b', r'utili[sz]ation', r'achievement', r'\bshare\b'],
    'amount':       [r'amount', r'salary', r'cost', r'value', r'\bpay\b', r'\bfee\b', r'fund', r'release',
                     r'price', r'\bbill\b', r'revenue', r'\btotal\b', r'\bbs\b', r'allowance', r'\bemolument\b'],
    'name':         [r'name', r'title'],
}

# Numeric ID-ish / non-metric names that must never be the main metric
NON_METRIC_PATTERNS = [r'\byear\b', r'\byr\b', r'^\s*fy\s*$', r'\bmonth\s*(no|number|#)?\b', r'\bage\b',
                       r'\bid\b', r'\bcode\b', r'\bno\.?$', r'\bcount\b', r'\bsr\.?\s*#?$', r'^#$', r'\bs\.?\s*#?\s*$',
                       r'\bserial\b', r'\bphone\b', r'\bmobile\b', r'\bcnic\b', r'\bpin\b', r'\bzip\b']

SEVERITY_CRITICAL = 'CRITICAL'
SEVERITY_WARNING = 'WARNING'
SEVERITY_INFO = 'INFO'

CATEGORICAL_ROLES = {'department', 'district', 'status', 'vendor', 'designation',
                     'grade', 'category', 'gender'}


# ------------------------------------------------------------------ #
# Helpers
# ------------------------------------------------------------------ #
def _matches(name, patterns):
    low = str(name).strip().lower().replace('_', ' ')
    return any(re.search(p, low) for p in patterns)


def detect_role(col_name):
    for role in ROLE_ORDER:
        if _matches(col_name, ROLE_PATTERNS[role]):
            return role
    return 'other'


def _j(v, nd=2):
    """JSON-safe number."""
    if v is None:
        return None
    try:
        if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
            return None
        f = float(v)
        if abs(f) >= 1e15:
            return f
        return round(f, nd)
    except (TypeError, ValueError):
        return None


def _is_numeric_dtype(s):
    return pd.api.types.is_numeric_dtype(s)


def _parse_numeric_string(val):
    """Parse 'Rs. 1,50,000', '(500)', '1.234.567,89', '$1,200.50' -> float or None."""
    if val is None or (isinstance(val, float) and np.isnan(val)):
        return None
    if isinstance(val, (int, float, np.integer, np.floating)):
        return float(val)
    s = str(val).strip()
    if s == '' or s.lower() in ERROR_STRINGS:
        return None
    # multiplier words ("2.5 million") are unreliable to convert - refuse rather than guess
    if re.search(r'\b(million|billion|crore|lakh|lac|thousand)\b', s.lower()):
        return None
    # alphanumeric codes like "BC-101" / "EMP-1005" must never become numbers
    if re.search(r'[A-Za-z]\s*-\s*\d', s):
        return None
    neg = False
    m = re.fullmatch(r'\(([^()]*)\)', s)
    if m:
        neg, s = True, m.group(1)
    s = re.sub(r'[^0-9.,\-]', '', s)
    if s in ('', '-', '.', ',', '--', '.-', '-.', '-,', '-'):
        return None
    sign = 1
    if '-' in s:
        if s.count('-') == 1 and s.index('-') <= 1:
            sign = -1
            s = s.replace('-', '')
        else:
            return None
    if s in ('', '.', ','):
        return None
    # currency-style leading dot (from "Rs." etc.) must be dropped BEFORE separator logic
    if s[0] in '.,':
        rest = s[1:].lstrip('.,')
        s = rest if (len(rest) > 3 or rest == '') else ('0' + s)
    if s == '':
        return None
    has_c, has_d = ',' in s, '.' in s
    if has_c and has_d:
        if re.fullmatch(r'\d{1,3}(,\d{2,3})+(\.\d+)?', s):
            s = s.replace(',', '')                     # comma = thousands
        else:
            s = s.replace('.', '').replace(',', '.')   # dot = thousands, comma = decimal
    elif has_c:
        if (re.fullmatch(r'\d{1,3}(,\d{3})+', s) or re.fullmatch(r'\d{1,2}(,\d{2})+(,\d{3})', s)
                or re.fullmatch(r'\d+,\d{3}', s)):
            s = s.replace(',', '')                     # thousands separator (incl. Indian grouping)
        else:
            s = s.replace(',', '.')                    # decimal comma
    elif has_d and s.count('.') > 1:
        head, _, tail = s.rpartition('.')
        s = head.replace('.', '') + '.' + tail         # keep only last dot
    s = s.rstrip('.')
    if s == '':
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if (neg or sign < 0) else v


def _coerce_numeric(series):
    """Try to convert an object column to numeric. Returns (new_series, rate)."""
    nonnull = series.dropna()
    if len(nonnull) == 0:
        return series, 0.0
    parsed = nonnull.map(_parse_numeric_string)
    rate = parsed.notna().sum() / len(nonnull)
    if rate < 0.85:
        return series, rate
    new = series.map(_parse_numeric_string)
    return new, rate


def _parse_dates(series, prefer_dayfirst=True):
    """Try both day-first and month-first parsing; return best (parsed_series, rate, dayfirst)."""
    s = series.dropna()
    if len(s) == 0:
        return series, 0.0, prefer_dayfirst
    best, best_rate, best_df = None, -1.0, prefer_dayfirst
    for df_flag in ([True, False] if prefer_dayfirst else [False, True]):
        try:
            parsed = pd.to_datetime(s, errors='coerce', dayfirst=df_flag)
        except (ValueError, TypeError, OverflowError):
            continue
        # keep only plausible dates
        valid = parsed.dropna()
        plausible = valid[(valid.dt.year >= 1950) & (valid.dt.year <= 2100)]
        rate = len(plausible) / len(s)
        if rate > best_rate + 1e-9:
            best, best_rate, best_df = parsed, rate, df_flag
        if best_rate >= 0.999:
            break
    if best is None:
        return series, 0.0, prefer_dayfirst
    out = series.copy()
    out.loc[s.index] = best
    # invalidate implausible years
    try:
        bad = out.notna() & (~out.dt.year.between(1950, 2100))
        out[bad] = pd.NaT
    except (TypeError, AttributeError):
        pass
    return out, best_rate, best_df


# ------------------------------------------------------------------ #
# File loading
# ------------------------------------------------------------------ #
def _looks_headerless(df):
    """Header heuristics: mostly 'Unnamed' columns, or ALL column names numeric."""
    if len(df.columns) == 0 or len(df) == 0:
        return False
    unnamed = sum(1 for c in df.columns if str(c).startswith('Unnamed'))
    if unnamed / len(df.columns) >= 0.6:
        return True
    try:
        if all(pd.api.types.is_number(c) or str(c).strip().isdigit() for c in df.columns):
            return True
    except Exception:
        pass
    return False


def _read_sheet_excel(path, sheet):
    df = pd.read_excel(path, sheet_name=sheet, header=0)
    if _looks_headerless(df):
        # likely no header row -> generic names
        df = pd.read_excel(path, sheet_name=sheet, header=None)
        df.columns = [f'Column_{i+1}' for i in range(len(df.columns))]
        df.attrs['no_header'] = True
    return df


def load_file(path, filename=''):
    """Load workbook/CSV. Returns dict: sheets -> {name: df}, meta info."""
    ext = (filename or path).lower().rsplit('.', 1)[-1]
    sheets, notes, wb_meta = {}, [], {'hidden_sheets': [], 'merged_ranges': {}}
    if ext == 'csv':
        import csv as _csv
        # deterministic delimiter detection from the first data line
        sep = ','
        try:
            with open(path, 'r', encoding='utf-8', errors='replace') as fh:
                for line in fh:
                    if line.strip():
                        counts = {d: line.count(d) for d in (',', ';', '\t', '|')}
                        best = max(counts, key=counts.get)
                        if counts[best] > 0:
                            sep = best
                        break
        except Exception:
            pass
        df, last_err = None, None
        for enc in ('utf-8-sig', 'utf-8', 'latin-1'):
            try:
                df = pd.read_csv(path, encoding=enc, sep=sep)
                break
            except pd.errors.EmptyDataError:
                raise ValueError('The CSV file is empty — no data found.')
            except (UnicodeDecodeError, pd.errors.ParserError, _csv.Error, ValueError) as e:
                last_err = e
                continue
        if df is None:
            raise ValueError(f'Could not read this CSV file ({type(last_err).__name__} — encoding not recognised). '
                             'Please re-save it as XLSX and try again.')
        if _looks_headerless(df):
            try:
                df = pd.read_csv(path, encoding=enc, sep=sep, header=None)
                df.columns = [f'Column_{i+1}' for i in range(len(df.columns))]
                notes.append('Header row not detected in CSV — generic column names assigned (first row treated as data).')
            except Exception:
                pass  # keep the original interpretation
        sheets['CSV_Data'] = df
    elif ext in ('xlsx', 'xlsm', 'xls'):
        try:
            all_sheets = pd.read_excel(path, sheet_name=None, header=0)
        except Exception as e:
            raise ValueError(f'Could not read this Excel file: {e}')
        # workbook-level meta (hidden sheets, merged cells) best-effort
        if ext in ('xlsx', 'xlsm'):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
                wb_meta['hidden_sheets'] = [ws.title for ws in wb.worksheets if ws.sheet_state != 'visible']
                for ws in wb.worksheets:
                    if ws.merged_cells.ranges:
                        wb_meta['merged_ranges'][ws.title] = len(ws.merged_cells.ranges)
                wb.close()
            except Exception:
                pass
        for name, raw in all_sheets.items():
            df = raw
            if _looks_headerless(df):
                try:
                    df = pd.read_excel(path, sheet_name=name, header=None)
                    df.columns = [f'Column_{i+1}' for i in range(len(df.columns))]
                    notes.append(f"Sheet '{name}': header row not detected — generic column names assigned (first row treated as data).")
                except Exception:
                    df = raw
            sheets[name] = df
    else:
        raise ValueError(f'Unsupported file type ".{ext}". Please upload .xlsx, .xls, .xlsm or .csv')
    sheets = {k: v for k, v in sheets.items()}
    return sheets, notes, wb_meta


# ------------------------------------------------------------------ #
# Main analysis pipeline
# ------------------------------------------------------------------ #
def analyze(raw_df, options=None, context=None):
    """
    options: {standardize_case: bool, fill_missing: bool, remove_duplicates: bool}
    context: {file_name, sheet_name, sheet_notes, hidden_sheets, merged_ranges, sheets_profile}
    Returns result dict (contains DataFrames + payload-safe values).
    """
    options = options or {}
    ctx = context or {}
    raw = raw_df.copy()
    raw.columns = [str(c).strip() for c in raw.columns]
    n_rows, n_cols = raw.shape

    issues, change_log, assumptions = [], [], []

    # ---------------- 1. column metadata ----------------
    col_meta = {}
    for c in raw.columns:
        s = raw[c]
        role = detect_role(c)
        is_num = _is_numeric_dtype(s)
        is_date = pd.api.types.is_datetime64_any_dtype(s)
        col_meta[c] = {
            'role': role, 'dtype': str(s.dtype),
            'is_numeric': is_num, 'is_datetime': is_date,
            'is_object': s.dtype == object,
        }

    # treat empty strings / error strings as missing for profiling
    norm = raw.copy()
    for c in norm.columns:
        if norm[c].dtype == object:
            norm[c] = norm[c].map(lambda v: np.nan if (isinstance(v, str) and (v.strip() == '' or v.strip().lower() in ERROR_STRINGS)) else v)

    missing_by_col = norm.isna().sum()

    # ---------------- 2. profiling ----------------
    columns_profile = []
    for c in raw.columns:
        m = col_meta[c]
        s = norm[c]
        nonnull = s.dropna()
        uniq = s.nunique(dropna=True)
        entry = {
            'name': c, 'dtype': m['dtype'], 'role': m['role'],
            'unique': int(uniq), 'missing': int(s.isna().sum()),
            'missing_pct': _j(s.isna().sum() / n_rows * 100, 1) if n_rows else 0,
            'min': None, 'max': None, 'sample': [],
        }
        if m['is_numeric'] and len(nonnull):
            entry['min'] = _j(nonnull.min()); entry['max'] = _j(nonnull.max())
        elif m['is_datetime'] and len(nonnull):
            entry['min'] = str(nonnull.min().date()); entry['max'] = str(nonnull.max().date())
        elif len(nonnull):
            entry['sample'] = [str(x)[:40] for x in nonnull.value_counts().head(3).index.tolist()]
        columns_profile.append(entry)

    # blank rows / columns
    blank_rows = int(norm.isna().all(axis=1).sum())
    blank_cols = [c for c in raw.columns if norm[c].isna().all()]
    if blank_rows:
        issues.append({'severity': SEVERITY_WARNING, 'area': 'Structure',
                       'description': f'{blank_rows} completely blank row(s) found.', 'count': blank_rows})
    if blank_cols:
        issues.append({'severity': SEVERITY_WARNING, 'area': 'Structure',
                       'description': f'Completely blank column(s): {", ".join(blank_cols)}', 'count': len(blank_cols)})

    # error-string cells
    err_cells = {}
    for c in raw.columns:
        if raw[c].dtype == object:
            cnt = raw[c].map(lambda v: isinstance(v, str) and v.strip().lower() in ERROR_STRINGS).sum()
            if cnt:
                err_cells[c] = int(cnt)
    if err_cells:
        tot = sum(err_cells.values())
        issues.append({'severity': SEVERITY_CRITICAL, 'area': 'Error values',
                       'description': f'Formula/error strings found ({tot} cells): ' +
                                      ', '.join(f'{k}: {v}' for k, v in err_cells.items()), 'count': tot})

    # ---------------- 3. duplicates ----------------
    dup_mask = norm.duplicated(keep='first')
    dup_rows = int(dup_mask.sum())
    dup_sample = []
    if dup_rows:
        idxs = norm[dup_mask].index[:10].tolist()
        dup_sample = idxs
        issues.append({'severity': SEVERITY_CRITICAL, 'area': 'Duplicates',
                       'description': f'{dup_rows} exact duplicate record(s) detected (all columns identical).',
                       'count': dup_rows})

    id_cols = [c for c in raw.columns if col_meta[c]['role'] == 'id']
    id_dup_info = None
    if id_cols:
        for idc in id_cols:
            s = norm[idc].dropna()
            dupd = s[s.duplicated(keep=False)]
            if len(dupd):
                vc = dupd.value_counts()
                id_dup_info = {
                    'column': idc,
                    'top': [{'value': str(k), 'times': int(v)} for k, v in vc.head(10).items()],
                    'count': int(vc.sum()),
                    'distinct': int(dupd.nunique()),
                }
                issues.append({'severity': SEVERITY_CRITICAL, 'area': 'Duplicates',
                               'description': f'Column "{idc}" has {dupd.nunique()} value(s) appearing more than once '
                                              f'({int(dupd.value_counts().sum())} rows involved).',
                               'count': int(dupd.value_counts().sum())})
                break

    # ---------------- 4. cleaning ----------------
    cleaned = raw.copy()

    # 4a. error strings -> NaN (logged)
    for c in cleaned.columns:
        if cleaned[c].dtype == object:
            mask = cleaned[c].map(lambda v: isinstance(v, str) and v.strip().lower() in ERROR_STRINGS)
            if mask.sum():
                change_log.append({'field': c, 'original': 'error value', 'updated': '(blank)',
                                   'reason': 'Excel error / N-A string converted to blank', 'cells': int(mask.sum())})
                cleaned.loc[mask, c] = np.nan

    # 4b. trim + collapse whitespace
    for c in cleaned.columns:
        if cleaned[c].dtype == object:
            def _clean_txt(v):
                if isinstance(v, str):
                    nv = re.sub(r'\s+', ' ', v.strip())
                    return nv if nv != '' else np.nan
                return v
            before = cleaned[c].copy()
            cleaned[c] = cleaned[c].map(_clean_txt)
            changed = int((before.fillna('__NA__') != cleaned[c].fillna('__NA__')).sum())
            if changed:
                issues.append({'severity': SEVERITY_INFO, 'area': 'Text formatting',
                               'description': f'Leading/trailing/extra spaces cleaned in "{c}" ({changed} cells).',
                               'count': changed})
                change_log.append({'field': c, 'original': 'values with extra spaces',
                                   'updated': 'trimmed values', 'reason': 'Whitespace clean-up (TRIM/CLEAN)',
                                   'cells': changed})

    # 4c. numeric coercion of amount-like object columns (never ID/code columns)
    coerced_numeric_cols = []
    for c in cleaned.columns:
        m = col_meta[c]
        if m['is_object'] and m['role'] in ('amount', 'budget', 'expenditure') \
                and not _matches(c, [r'\bcode\b', r'\bid\b', r'\bno\.?$', r'\bref\b', r'\bkey\b', r'\bserial\b', r'\bnumber\b']):
            new, rate = _coerce_numeric(cleaned[c])
            if rate >= 0.85 and new.notna().sum() > 0 and not _is_numeric_dtype(cleaned[c]):
                # capture a real before/after example for the audit trail
                example = None
                for orig, conv in zip(cleaned[c].dropna().head(50), new.dropna().head(50)):
                    if isinstance(orig, str) and not _matches(orig, [r'^\s*-?\d+(\.\d+)?\s*$']):
                        example = (orig, conv)
                        break
                converted = int(new.notna().sum() - cleaned[c].isna().sum())
                change_log.append({'field': c,
                                   'original': f'text amounts{f" (e.g. \"{example[0]}\" -> {example[1]:g})" if example else ""}',
                                   'updated': 'numeric values',
                                   'reason': 'Converted text amounts to numbers for calculation (thousands separators / currency prefixes removed)',
                                   'cells': max(converted, 1)})
                col_meta[c]['is_numeric'] = True
                coerced_numeric_cols.append(c)
                cleaned[c] = new

    # 4d. date parsing
    date_cols, primary_date_col = {}, None
    for c in cleaned.columns:
        m = col_meta[c]
        if m['is_datetime']:
            date_cols[c] = {'rate': 1.0, 'dayfirst': None}
            primary_date_col = primary_date_col or c
            continue
        if m['is_object'] and m['role'] == 'date':
            parsed, rate, dayfirst = _parse_dates(cleaned[c])
            if rate >= 0.5:
                unparsed = int(parsed.isna().sum() - cleaned[c].isna().sum())
                change_log.append({'field': c, 'original': 'mixed text dates',
                                   'updated': 'standardised dates (dd-mmm-yyyy)',
                                   'reason': f'Date standardisation (day-first={dayfirst})', 'cells': int(rate * len(cleaned[c].dropna()))})
                if unparsed > 0:
                    issues.append({'severity': SEVERITY_WARNING, 'area': 'Dates',
                                   'description': f'{unparsed} value(s) in "{c}" could not be interpreted as dates.',
                                   'count': unparsed})
                cleaned[c] = parsed
                col_meta[c]['is_datetime'] = True
                date_cols[c] = {'rate': rate, 'dayfirst': dayfirst}
                primary_date_col = primary_date_col or c
    if date_cols and any(v['dayfirst'] for v in date_cols.values()):
        assumptions.append('Ambiguous date values were interpreted day-first (dd/mm/yyyy), the usual convention in Pakistani records.')

    # 4e. case standardisation for categorical columns
    if options.get('standardize_case', True):
        target_cat_cols = [c for c in cleaned.columns
                           if col_meta[c]['role'] in CATEGORICAL_ROLES and cleaned[c].dtype == object]
        for c in target_cat_cols:
            s = cleaned[c].dropna()
            if len(s) == 0:
                continue
            groups = {}
            for v in s:
                groups.setdefault(str(v).strip().lower(), {}).setdefault(str(v).strip(), 0)
                groups[str(v).strip().lower()][str(v).strip()] += 1
            mapping, changed_cells, pairs = {}, 0, []
            for key, variants in groups.items():
                if len(variants) <= 1:
                    continue
                # canonical: prefer a properly Title Case spelling, then the most frequent
                def _score(var_cnt):
                    var, cnt = var_cnt
                    return (-1 if var == var.title() else 0, -cnt, len(var))
                canonical = sorted(variants.items(), key=_score)[0][0]
                for var, cnt in variants.items():
                    if var != canonical:
                        mapping[var] = canonical
                        changed_cells += cnt
                        pairs.append((var, canonical, cnt))
            if mapping:
                mask = cleaned[c].isin(mapping)
                for var, canonical, cnt in pairs[:12]:
                    change_log.append({'field': c, 'original': var, 'updated': canonical,
                                       'reason': 'Case/spacing standardisation', 'cells': cnt})
                issues.append({'severity': SEVERITY_INFO, 'area': 'Standardisation',
                               'description': f'"{c}": {len(pairs)} label variant(s) unified to a single consistent spelling '
                                              f'(e.g. "{pairs[0][0]}" -> "{pairs[0][1]}").', 'count': changed_cells})
                cleaned.loc[mask, c] = cleaned[c].map(lambda v: mapping.get(v, v))

    # 4f. missing-value marking
    if options.get('fill_missing', True):
        for c in cleaned.columns:
            if col_meta[c]['role'] in CATEGORICAL_ROLES and cleaned[c].dtype == object:
                cnt = int(cleaned[c].isna().sum())
                if cnt:
                    change_log.append({'field': c, 'original': '(blank)', 'updated': 'MISSING',
                                       'reason': 'Blank category clearly marked (no value invented)', 'cells': cnt})
                    cleaned[c] = cleaned[c].fillna('MISSING')

    # 4g. duplicate removal (optional, default OFF)
    removed_dup = 0
    kept_orig_idx = list(range(n_rows))
    dropped_orig = []
    if options.get('remove_duplicates', False) and dup_rows:
        keep_mask = ~cleaned.duplicated(keep='first')
        removed_dup = int((~keep_mask).sum())
        kept_orig_idx = [int(i) for i in cleaned.index[keep_mask]]
        dropped_orig = [int(i) for i in cleaned.index[~keep_mask]]
        cleaned = cleaned[keep_mask].reset_index(drop=True)
        change_log.append({'field': '(all columns)', 'original': f'{removed_dup} duplicate row(s)',
                           'updated': 'removed', 'reason': 'User-selected option: delete exact duplicates', 'cells': removed_dup})
    elif dup_rows:
        assumptions.append(f'{dup_rows} exact duplicate row(s) were NOT deleted (kept for audit). '
                           'See Exception_Report; enable "Remove exact duplicates" to drop them from Cleaned_Data.')

    # ---------------- 5. missing-data classification ----------------
    crit_roles = {'id', 'name', 'date', 'amount', 'budget', 'expenditure'}
    for c in raw.columns:
        miss = int(missing_by_col[c])
        if miss == 0:
            continue
        role = col_meta[c]['role']
        sev = SEVERITY_CRITICAL if role in crit_roles else (SEVERITY_WARNING if role in CATEGORICAL_ROLES else SEVERITY_INFO)
        issues.append({'severity': sev, 'area': 'Missing data',
                       'description': f'"{c}" has {miss} missing value(s) ({missing_by_col[c]/n_rows*100:.1f}%). '
                                      f'Classification: {"critical" if sev==SEVERITY_CRITICAL else ("important" if sev==SEVERITY_WARNING else "optional")} field.',
                       'count': miss})
    total_missing = int(missing_by_col.sum())

    # ---------------- 6. metric & dimension selection ----------------
    numeric_cols = [c for c in cleaned.columns if _is_numeric_dtype(cleaned[c]) and not cleaned[c].isna().all()]
    metric_candidates = [c for c in numeric_cols
                         if col_meta[c]['role'] in ('amount', 'budget', 'expenditure')
                         and not _matches(c, NON_METRIC_PATTERNS)]
    other_numeric = [c for c in numeric_cols
                     if c not in metric_candidates and not _matches(c, NON_METRIC_PATTERNS)
                     and col_meta[c]['role'] not in ('id', 'percent')]
    main_metric = None
    _role_pref = {'amount': 0, 'expenditure': 1, 'budget': 2}
    if metric_candidates:
        main_metric = sorted(metric_candidates, key=lambda c: (_role_pref.get(col_meta[c]['role'], 3),
                                                              -abs(cleaned[c].sum(skipna=True))))[0]
    elif other_numeric:
        main_metric = max(other_numeric, key=lambda c: cleaned[c].abs().sum(skipna=True))
    secondary_metric = None
    pool = [c for c in metric_candidates + other_numeric if c != main_metric]
    if pool:
        secondary_metric = max(pool, key=lambda c: cleaned[c].abs().sum(skipna=True))

    dim_priority = ['department', 'district', 'status', 'category', 'grade', 'designation', 'vendor', 'gender']
    dim_cols = []
    for role in dim_priority:
        for c in cleaned.columns:
            if col_meta[c]['role'] == role and cleaned[c].dtype == object and cleaned[c].nunique(dropna=True) >= 1:
                dim_cols.append(c)
                break
    if not dim_cols and len(cleaned):
        for c in cleaned.columns:
            if cleaned[c].dtype == object and 1 < cleaned[c].nunique(dropna=True) <= min(50, max(10, int(n_rows * 0.05))):
                dim_cols.append(c)
                if len(dim_cols) >= 2:
                    break
    dim_cols = dim_cols[:4]

    # constant columns
    for c in numeric_cols:
        if cleaned[c].nunique(dropna=True) <= 1:
            issues.append({'severity': SEVERITY_INFO, 'area': 'Structure',
                           'description': f'Column "{c}" has a single constant value for all rows.', 'count': n_rows})

    # ---------------- 7. negative & percent checks ----------------
    for c in numeric_cols:
        if col_meta[c]['role'] in ('amount', 'budget', 'expenditure'):
            neg = int((cleaned[c] < 0).sum())
            if neg:
                issues.append({'severity': SEVERITY_WARNING, 'area': 'Values',
                               'description': f'{neg} negative value(s) in "{c}" - retained as-is; may require review.',
                               'count': neg})
    for c in numeric_cols:
        if col_meta[c]['role'] == 'percent' or _matches(c, [r'percent', r'%']):
            bad = int(((cleaned[c] < 0) | (cleaned[c] > 100)).sum())
            if bad:
                issues.append({'severity': SEVERITY_WARNING, 'area': 'Values',
                               'description': f'{bad} value(s) in "{c}" fall outside the logical 0-100% range.',
                               'count': bad})

    # ---------------- 8. outliers (IQR) ----------------
    outliers, outlier_flags = [], {}
    for c in numeric_cols:
        s = cleaned[c].dropna()
        if len(s) < 10 or s.nunique() < 5:
            continue
        q1, q3 = s.quantile(0.25), s.quantile(0.75)
        iqr = q3 - q1
        if iqr <= 0:
            continue
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        mask = (cleaned[c] < lo) | (cleaned[c] > hi)
        cnt = int(mask.sum())
        if cnt:
            outlier_flags[c] = (lo, hi, cnt)
            for idx in cleaned[c][mask].index[:25]:
                val = cleaned.loc[idx, c]
                orig_row = kept_orig_idx[int(idx)] if int(idx) < len(kept_orig_idx) else int(idx)
                outliers.append({
                    'type': 'Statistical anomaly (IQR)',
                    'row': orig_row + 2,  # +2 = header + 1-based, refers to Raw_Data
                    'column': c, 'value': _j(val),
                    'detail': f'Outside typical range {_j(lo)} to {_j(hi)} - requires review (not proof of error).',
                })
    total_outlier_rows = len(outliers)
    if outlier_flags:
        issues.append({'severity': SEVERITY_INFO, 'area': 'Anomalies',
                       'description': 'Potential anomalies (IQR method): ' +
                                      ', '.join(f'{k}: {v[2]}' for k, v in outlier_flags.items()) +
                                      '. Flagged for review only - no judgement implied.',
                       'count': sum(v[2] for v in outlier_flags.values())})

    # ---------------- 9. budget vs expenditure ----------------
    budget_col = next((c for c in cleaned.columns if col_meta[c]['role'] == 'budget' and _is_numeric_dtype(cleaned[c])), None)
    expend_col = next((c for c in cleaned.columns if col_meta[c]['role'] == 'expenditure' and _is_numeric_dtype(cleaned[c])), None)
    budget = None
    if budget_col and expend_col and budget_col != expend_col:
        tot_b = float(cleaned[budget_col].sum(skipna=True))
        tot_e = float(cleaned[expend_col].sum(skipna=True))
        dept_c = next((c for c in dim_cols if col_meta[c]['role'] == 'department'), None)
        per_dept = []
        if dept_c:
            g = cleaned.groupby(dept_c, dropna=False).agg(**{
                'budget': (budget_col, 'sum'), 'expenditure': (expend_col, 'sum'), 'records': (budget_col, 'count')})
            for label, r in g.iterrows():
                util = (r['expenditure'] / r['budget']) if r['budget'] else None
                per_dept.append({'label': str(label), 'budget': _j(r['budget']),
                                 'expenditure': _j(r['expenditure']), 'remaining': _j(r['budget'] - r['expenditure']),
                                 'utilization': _j(util * 100, 1) if util is not None else None,
                                 'records': int(r['records'])})
            per_dept.sort(key=lambda x: -(x['budget'] or 0))
        above = cleaned[cleaned[expend_col] > cleaned[budget_col]]
        above_rows = []
        for idx, r in above.head(20).iterrows():
            orig_row = kept_orig_idx[int(idx)] if int(idx) < len(kept_orig_idx) else int(idx)
            above_rows.append({'row': orig_row + 2, 'budget': _j(r[budget_col]),
                               'expenditure': _j(r[expend_col]),
                               'detail': f'Expenditure exceeds budget allocation (variance {_j(r[budget_col] - r[expend_col])}) - requires review.'})
        if len(above):
            issues.append({'severity': SEVERITY_WARNING, 'area': 'Budget',
                           'description': f'{len(above)} record(s) where expenditure exceeds the budget allocation - flagged for review.',
                           'count': int(len(above))})
        budget = {
            'budget_col': budget_col, 'expenditure_col': expend_col,
            'total_budget': _j(tot_b), 'total_expenditure': _j(tot_e),
            'remaining': _j(tot_b - tot_e),
            'utilization': _j(tot_e / tot_b * 100, 1) if tot_b else None,
            'per_dept': per_dept, 'above_budget_count': int(len(above)), 'above_budget_rows': above_rows,
        }
        if budget_col:
            assumptions.append(f'"{budget_col}" treated as budget allocation and "{expend_col}" as expenditure, based on column names. '
                               'Budget / release / commitment / actual payment are NOT assumed to be interchangeable.')
        issues.append({'severity': SEVERITY_INFO, 'area': 'Budget',
                       'description': f'Overall utilisation: {_j(tot_e / tot_b * 100, 1) if tot_b else "N/A"}% '
                                      f'(expenditure {_j(tot_e)} of budget {_j(tot_b)}).', 'count': n_rows})

    # ---------------- 10. descriptive statistics ----------------
    stats = []
    for c in (metric_candidates + other_numeric)[:10] or numeric_cols[:10]:
        s = cleaned[c].dropna()
        if len(s) == 0:
            continue
        stats.append({'column': c, 'count': int(len(s)), 'sum': _j(s.sum()), 'mean': _j(s.mean()),
                      'median': _j(s.median()), 'min': _j(s.min()), 'max': _j(s.max()),
                      'std': _j(s.std()) if len(s) > 1 else None})

    # ---------------- 11. pivot-style summaries ----------------
    pivots = []
    for d in dim_cols:
        s = cleaned[d].dropna()
        if len(s) == 0 or s.nunique() < 1:
            continue
        if main_metric and _is_numeric_dtype(cleaned[main_metric]):
            g = cleaned.groupby(d, dropna=False).agg(**{
                'count': (d, 'size'), 'sum': (main_metric, 'sum'), 'avg': (main_metric, 'mean'),
                'mn': (main_metric, 'min'), 'mx': (main_metric, 'max'), 'med': (main_metric, 'median')})
            g = g.sort_values('sum', ascending=False)
            total = g['sum'].sum()
            rows = [{'label': str(k), 'count': int(r['count']), 'sum': _j(r['sum']), 'avg': _j(r['avg']),
                     'min': _j(r['mn']), 'max': _j(r['mx']), 'median': _j(r['med']),
                     'pct': _j(r['sum'] / total * 100, 1) if total else None}
                    for k, r in g.head(15).iterrows()]
            pivots.append({'dimension': d, 'measure': main_metric, 'has_measure': True, 'rows': rows,
                           'total_labels': int(g.shape[0])})
        else:
            vc = cleaned[d].value_counts(dropna=False)
            rows = [{'label': str(k), 'count': int(v), 'pct': _j(v / n_rows * 100, 1)}
                    for k, v in vc.head(15).items()]
            pivots.append({'dimension': d, 'measure': None, 'has_measure': False, 'rows': rows,
                           'total_labels': int(vc.shape[0])})

    # status distribution (for donut)
    status_dist = None
    status_col = next((c for c in cleaned.columns if col_meta[c]['role'] == 'status'), None)
    if status_col:
        vc = cleaned[status_col].value_counts(dropna=False)
        status_dist = {'column': status_col,
                       'rows': [{'label': str(k), 'count': int(v)} for k, v in vc.head(8).items()]}

    # ---------------- 12. monthly trend ----------------
    trend = None
    if primary_date_col:
        s = cleaned[[primary_date_col]].dropna()
        if len(s) >= 3 and main_metric:
            tmp = cleaned[[primary_date_col, main_metric]].dropna()
            tmp['ym'] = pd.to_datetime(tmp[primary_date_col]).dt.to_period('M')
            g = tmp.groupby('ym').agg(values=(main_metric, 'sum'), counts=(main_metric, 'size')).sort_index()
            g = g.tail(36)
            trend = {'date_col': primary_date_col, 'measure': main_metric,
                     'labels': [str(p) for p in g.index],
                     'values': [_j(v) for v in g['values']],
                     'counts': [int(c) for c in g['counts']]}
        elif len(s) >= 3:
            tmp = cleaned[[primary_date_col]].dropna()
            tmp['ym'] = pd.to_datetime(tmp[primary_date_col]).dt.to_period('M')
            g = tmp.groupby('ym').size().sort_index().tail(36)
            trend = {'date_col': primary_date_col, 'measure': None,
                     'labels': [str(p) for p in g.index],
                     'values': [int(v) for v in g], 'counts': [int(v) for v in g]}
        date_range = (str(pd.to_datetime(cleaned[primary_date_col].dropna()).min().date()),
                      str(pd.to_datetime(cleaned[primary_date_col].dropna()).max().date())) if primary_date_col in date_cols else None
    else:
        date_range = None
    if date_range:
        assumptions.append(f'"{primary_date_col}" used as the reference date column (range {date_range[0]} to {date_range[1]}).')

    # ---------------- 13. top / bottom rankings (top 50 kept) ----------------
    top50, bottom50, rank_info = None, None, None
    if main_metric:
        s = cleaned[cleaned[main_metric].notna()]
        idc = next((c for c in cleaned.columns if col_meta[c]['role'] == 'id'), None)
        namec = next((c for c in cleaned.columns if col_meta[c]['role'] == 'name' and cleaned[c].dtype == object), None)
        deptc = next((c for c in cleaned.columns if col_meta[c]['role'] == 'department'), None)
        show_cols = [c for c in [idc, namec, deptc] if c] + ([primary_date_col] if primary_date_col else []) + \
                    [main_metric] + ([secondary_metric] if secondary_metric and secondary_metric not in (idc, namec, deptc) else [])
        ranked = s.sort_values(main_metric, ascending=False).head(50)
        def _rows(df):
            out = []
            for i, (idx, r) in enumerate(df.iterrows(), 1):
                row = {'rank': i}
                for c in show_cols:
                    v = r[c]
                    if isinstance(v, pd.Timestamp):
                        row[c] = str(v.date())
                    else:
                        row[c] = _j(v) if isinstance(v, (int, float, np.integer, np.floating)) else (str(v) if pd.notna(v) else '')
                out.append(row)
            return out
        top50 = {'columns': show_cols, 'rows': _rows(ranked)}
        bottom50 = {'columns': show_cols, 'rows': _rows(s.sort_values(main_metric, ascending=True).head(50))}
        rank_info = {'metric': main_metric, 'direction': 'descending (highest = rank 1)'}

    # ---------------- 14. duplicate sample rows ----------------
    dup_records = []
    if dup_rows:
        cols_show = [c for c in ([next((c for c in cleaned.columns if col_meta[c]['role'] == 'id'), None)] if id_cols else [])]
        sample_idx = dup_sample[:10]
        for idx in sample_idx:
            rec = {'row': int(idx) + 2}
            for c in raw.columns[:8]:
                v = raw.loc[idx, c]
                rec[c] = str(v) if not isinstance(v, float) or not np.isnan(v) else ''
            dup_records.append(rec)

    # ---------------- 15. calculations (row-level audit) ----------------
    calc_data = pd.DataFrame({'Row (Raw_Data)': list(range(2, n_rows + 2))})
    if removed_dup:
        dropped_set = set(dropped_orig)
        calc_data['In Cleaned_Data'] = ['REMOVED (exact duplicate)' if i in dropped_set else 'KEPT' for i in range(n_rows)]
    if dup_rows:
        dup_set = set(int(i) for i in norm[dup_mask].index)
        calc_data['Duplicate in Raw'] = ['YES' if i in dup_set else '' for i in range(n_rows)]
    miss_row = raw.isna().sum(axis=1)
    if n_rows and len(raw.columns):
        miss_row = miss_row + raw.eq('').sum(axis=1).astype(int)
    calc_data['Missing fields'] = miss_row.values
    if main_metric and _is_numeric_dtype(cleaned[main_metric]):
        rank_series = cleaned[main_metric].rank(ascending=False, method='min')
        rank_map = {orig: rk for orig, rk in zip(kept_orig_idx, rank_series.values)}
        calc_data[f'Rank ({main_metric})'] = [rank_map.get(i) for i in range(n_rows)]
    if outlier_flags:
        oflag = [''] * n_rows
        for o in outliers:
            oi = o['row'] - 2
            if 0 <= oi < n_rows:
                oflag[oi] = (oflag[oi] + '; ' if oflag[oi] else '') + o['column']
        calc_data['Outlier flag'] = oflag

    # ---------------- 16. KPIs ----------------
    top_dept = None
    if dim_cols:
        d0 = dim_cols[0]
        vc = cleaned[d0].value_counts(dropna=False)
        if len(vc):
            top_dept = {'column': d0, 'label': str(vc.index[0]), 'count': int(vc.iloc[0])}

    kpis = {
        'total_records': int(n_rows), 'columns': int(n_cols),
        'duplicate_rows': dup_rows, 'missing_cells': total_missing,
        'missing_pct': _j(total_missing / (n_rows * n_cols) * 100, 1) if n_rows and n_cols else 0,
        'outliers': sum(v[2] for v in outlier_flags.values()),
        'date_range': list(date_range) if date_range else None,
        'main_metric': main_metric,
        'main_metric_total': _j(cleaned[main_metric].sum(skipna=True)) if main_metric and _is_numeric_dtype(cleaned[main_metric]) else None,
        'top_category': top_dept,
        'utilization': budget['utilization'] if budget else None,
        'total_budget': budget['total_budget'] if budget else None,
        'total_expenditure': budget['total_expenditure'] if budget else None,
    }

    # ---------------- 17. assumptions & notes ----------------
    if ctx.get('sheet_name'):
        assumptions.insert(0, f"Sheet \"{ctx['sheet_name']}\" was selected for analysis "
                              f"({n_rows} rows x {n_cols} columns) - it is the largest data sheet in the workbook."
                              if ctx.get('auto_selected') else
                              f'Sheet "{ctx["sheet_name"]}" selected by user ({n_rows} rows x {n_cols} columns).')
    if ctx.get('sheet_notes'):
        assumptions.extend(ctx['sheet_notes'])
    assumptions.append('Only whitespace / case standardisation was applied to text. Distinct category labels were never merged and no values were invented.')
    assumptions.append('Formula cells are read as their last saved (cached) values.')
    if budget is None and main_metric and col_meta.get(main_metric, {}).get('role') in ('amount',):
        assumptions.append(f'"{main_metric}" treated as the main amount measure based on its name; the file does not state the currency unit, so figures are shown as-is.')
    if not dup_rows and not total_missing:
        pass
    if removed_dup:
        assumptions.append(f'{removed_dup} exact duplicate row(s) removed from Cleaned_Data on user request. Raw_Data retains every original row.')

    # ---------------- 18. validation rules & power query recipe ----------------
    validation_rules = build_validation_rules(cleaned, col_meta)
    powerquery = generate_powerquery(cleaned, col_meta, coerced_numeric_cols, options, ctx)
    kpis['sheet_name'] = ctx.get('sheet_name', '')
    mgmt_summary = build_management_summary(kpis, pivots, budget, issues, trend, top50)
    formulas = build_formula_recipes(raw, col_meta, ctx, main_metric, dim_cols, id_dup_info)

    # builder-friendly column lists
    categorical_columns = []
    for c in cleaned.columns:
        if cleaned[c].dtype == object:
            nun = int(cleaned[c].nunique(dropna=True))
            if 1 <= nun <= 100:
                categorical_columns.append({
                    'name': c, 'unique': nun,
                    'values': [str(v) for v in cleaned[c].value_counts(dropna=False).head(50).index.tolist()],
                })
    numeric_columns = [c for c in cleaned.columns if _is_numeric_dtype(cleaned[c]) and cleaned[c].notna().any()]

    # severity ordering
    sev_rank = {SEVERITY_CRITICAL: 0, SEVERITY_WARNING: 1, SEVERITY_INFO: 2}
    issues.sort(key=lambda x: sev_rank.get(x['severity'], 3))

    result = {
        'raw': raw, 'cleaned': cleaned, 'calc': calc_data,
        'col_meta': col_meta, 'columns_profile': columns_profile,
        'issues': issues, 'change_log': change_log, 'assumptions': assumptions,
        'kpis': kpis, 'stats': stats, 'pivots': pivots, 'status_dist': status_dist,
        'trend': trend, 'top50': top50, 'bottom50': bottom50, 'rank_info': rank_info,
        'outliers': outliers[:60], 'outlier_total': sum(v[2] for v in outlier_flags.values()),
        'budget': budget, 'dup_records': dup_records, 'dup_rows': dup_rows,
        'id_dup_info': id_dup_info, 'date_cols': list(date_cols.keys()),
        'primary_date_col': primary_date_col, 'main_metric': main_metric,
        'dim_cols': dim_cols, 'n_rows': n_rows, 'n_cols': n_cols,
        'options_applied': options, 'ctx': ctx,
        'validation_rules': validation_rules, 'powerquery': powerquery,
        'categorical_columns': categorical_columns, 'numeric_columns': numeric_columns,
        'coerced_numeric_cols': coerced_numeric_cols, 'mgmt_summary': mgmt_summary,
        'formulas': formulas,
    }
    return result


# ------------------------------------------------------------------ #
# Payload (JSON-safe) for the web UI
# ------------------------------------------------------------------ #
def to_payload(result, job_id, filename, sheet_name, sheets_info, download_url):
    id_dup = None
    if result.get('id_dup_info'):
        d = result['id_dup_info']
        id_dup = {'column': d['column'], 'count': d['count'], 'distinct': d['distinct'],
                  'top': d['top'][:5]}
    return {
        'job_id': job_id,
        'file_name': filename,
        'sheet_name': sheet_name,
        'sheets_info': sheets_info,
        'download_url': download_url,
        'kpis': result['kpis'],
        'columns': result['columns_profile'],
        'issues': result['issues'],
        'change_log': result['change_log'],
        'assumptions': result['assumptions'],
        'stats': result['stats'],
        'pivots': result['pivots'],
        'status_dist': result['status_dist'],
        'trend': result['trend'],
        'top50': result['top50'],
        'bottom50': result['bottom50'],
        'rank_info': result['rank_info'],
        'outliers': result['outliers'],
        'outlier_total': result['outlier_total'],
        'budget': result['budget'],
        'dup_records': result['dup_records'],
        'dup_rows': result['dup_rows'],
        'id_dup': id_dup,
        'options_applied': result['options_applied'],
        'validation_rules': result.get('validation_rules', []),
        'powerquery': result.get('powerquery', {}),
        'categorical_columns': result.get('categorical_columns', []),
        'numeric_columns': result.get('numeric_columns', []),
        'join_report': result.get('join_report'),
        'is_merged': bool(result.get('merged_data') is not None),
        'mgmt_summary': result.get('mgmt_summary', []),
        'lookup_report': result.get('lookup_report'),
        'formulas': result.get('formulas', []),
    }




# ================================================================== #
# PIVOT BUILDER - full pivot options (rows, columns, measure, agg)
# ================================================================== #
AGG_FUNCS = {'sum': 'sum', 'avg': 'mean', 'mean': 'mean', 'min': 'min', 'max': 'max', 'median': 'median'}
AGG_LABELS = {'sum': 'Sum', 'avg': 'Average', 'min': 'Minimum', 'max': 'Maximum', 'median': 'Median'}


def build_pivot(df, row_dim, col_dim=None, measure=None, agg='sum', filter_col=None, filter_val=None):
    """Interactive pivot: rows x optional columns, measure + aggregation, optional filter."""
    if row_dim not in df.columns:
        raise ValueError(f'Row field "{row_dim}" not found.')
    if col_dim and col_dim not in df.columns:
        raise ValueError(f'Column field "{col_dim}" not found.')
    count_mode = measure in (None, '', '__count__')
    if not count_mode and measure not in df.columns:
        raise ValueError(f'Value field "{measure}" not found.')
    if not count_mode and not _is_numeric_dtype(df[measure]):
        raise ValueError(f'Value field "{measure}" is not numeric - choose "Record count" or a numeric column.')
    agg = (agg or 'sum').lower()
    if not count_mode and agg not in AGG_FUNCS:
        raise ValueError(f'Unknown aggregation "{agg}".')

    d = df
    applied_filter = None
    if filter_col and filter_val not in (None, '', '__all__'):
        if filter_col not in d.columns:
            raise ValueError(f'Filter field "{filter_col}" not found.')
        d = d[d[filter_col].astype(str) == str(filter_val)]
        applied_filter = {'column': filter_col, 'value': str(filter_val)}
    if len(d) == 0:
        raise ValueError('No records match the selected filter.')

    # normalise dimension labels (blanks -> "(blank)")
    rd = d[row_dim].fillna('(blank)').astype(str)
    cd = d[col_dim].fillna('(blank)').astype(str) if col_dim else None

    if not col_dim:
        if count_mode:
            g = rd.value_counts().sort_values(ascending=False)
            total = float(g.sum())
            rows = [{'label': str(k), 'value': int(v), 'pct': _j(v / total * 100, 1) if total else None}
                    for k, v in g.head(60).items()]
            return {'mode': 'single', 'row_dim': row_dim, 'col_dim': None, 'measure': None,
                    'agg': 'count', 'rows': rows, 'grand_total': _j(total),
                    'records': int(len(d)), 'filter': applied_filter}
        func = AGG_FUNCS[agg]
        tmp = pd.DataFrame({row_dim: rd, measure: d[measure].values})
        g = tmp.groupby(row_dim)[measure].agg(func).sort_values(ascending=False)
        total = float(g.sum()) if func == 'sum' else None
        rows = [{'label': str(k), 'value': _j(v),
                 'pct': _j(v / total * 100, 1) if total else None}
                for k, v in g.head(60).items()]
        return {'mode': 'single', 'row_dim': row_dim, 'col_dim': None, 'measure': measure,
                'agg': agg, 'rows': rows, 'grand_total': _j(total),
                'records': int(len(d)), 'filter': applied_filter}

    # cross-tab
    agg_func = None if count_mode else AGG_FUNCS[agg]
    tmp = pd.DataFrame({row_dim: rd, col_dim: cd})
    if count_mode:
        pt = pd.crosstab(tmp[row_dim], tmp[col_dim]).astype(float)
    else:
        tmp[measure] = d[measure].values
        pt = pd.pivot_table(tmp, index=row_dim, columns=col_dim, values=measure,
                            aggfunc=agg_func, dropna=False)
    row_tot = pt.sum(axis=1, skipna=True)
    col_tot = pt.sum(axis=0, skipna=True)
    grand = float(pt.sum(skipna=True).sum()) if count_mode or agg_func == 'sum' else None
    # sort and cap
    pt = pt.loc[row_tot.sort_values(ascending=False).index]
    pt = pt[col_tot.sort_values(ascending=False).index]
    pt = pt.iloc[:30, :15]
    rt = row_tot[pt.index]
    ct = col_tot[pt.columns]
    return {'mode': 'cross', 'row_dim': row_dim, 'col_dim': col_dim,
            'measure': None if count_mode else measure,
            'agg': 'count' if count_mode else agg,
            'row_labels': [str(i) for i in pt.index],
            'col_labels': [str(c) for c in pt.columns],
            'matrix': [[_j(v) for v in pt.iloc[r]] for r in range(len(pt))],
            'row_totals': [_j(v) for v in rt],
            'col_totals': [_j(v) for v in ct],
            'grand_total': _j(grand) if grand is not None else None,
            'records': int(len(d)), 'filter': applied_filter}


# ================================================================== #
# DATA MODELING - relationships between sheets (Excel Data Model style)
# ================================================================== #
def detect_relationships(sheets, max_pairs=15):
    """Detect candidate join keys between sheets. Never merges - only reports."""
    names = [n for n, d in sheets.items() if isinstance(d, pd.DataFrame) and d.shape[0] > 0 and d.shape[1] > 0]
    if len(names) < 2:
        return []
    rels = []

    def value_set(df, col, cap=20000):
        s = df[col].dropna()
        if len(s) == 0:
            return set()
        if len(s) > cap:
            s = s.sample(cap, random_state=0)
        return set(s.astype(str).str.strip().str.lower())

    for li in range(len(names)):
        for ri in range(len(names)):
            if li == ri:
                continue
            ln, rn = names[li], names[ri]
            left, right = sheets[ln], sheets[rn]
            best_for_pair = []
            for rcol in list(right.columns)[:60]:
                rset = value_set(right, rcol)
                if not rset or len(rset) > 5000:
                    continue
                for lcol in list(left.columns)[:80]:
                    same_name = str(lcol).strip().lower() == str(rcol).strip().lower()
                    if not same_name:
                        # only consider differently-named columns when values actually overlap
                        lsample = list(value_set(left, lcol, cap=400))
                        if not lsample or not (set(lsample) & rset):
                            continue
                    lset = value_set(left, lcol)
                    if not lset:
                        continue
                    inter = lset & rset
                    if not inter:
                        continue
                    matched_left = len(inter) / len(lset)          # % of left values found in right
                    matched_right = len(inter) / len(rset)         # % of right values used by left
                    if matched_left < 0.3 and not same_name:
                        continue
                    lser = left[lcol].dropna().astype(str).str.strip().str.lower()
                    rser = right[rcol].dropna().astype(str).str.strip().str.lower()
                    best_for_pair.append({
                        'left_sheet': ln, 'left_key': str(lcol),
                        'right_sheet': rn, 'right_key': str(rcol),
                        'left_rows': int(len(left)), 'right_rows': int(len(right)),
                        'matched_left_pct': _j(matched_left * 100, 1),
                        'matched_right_pct': _j(matched_right * 100, 1),
                        'left_key_unique_pct': _j(lser.nunique() / len(lser) * 100, 1) if len(lser) else 0,
                        'right_key_unique_pct': _j(rser.nunique() / len(rser) * 100, 1) if len(rser) else 0,
                        'duplicate_left_keys': int(len(lser) - lser.nunique()),
                        'duplicate_right_keys': int(len(rser) - rser.nunique()),
                        'same_name': same_name,
                        'matched_values': len(inter),
                    })
            # best candidates per sheet-pair: same-name first, then by coverage
            best_for_pair.sort(key=lambda r: (not r['same_name'], -(r['matched_left_pct'] or 0)))
            rels.extend(best_for_pair[:2])
    # global sort: same-name + high coverage first, cap results
    rels.sort(key=lambda r: (not r['same_name'], -((r['matched_left_pct'] or 0) * (r['matched_right_pct'] or 0))))
    # dedupe identical (reversed duplicates are kept as they describe different directions)
    seen, out = set(), []
    for r in rels:
        k = (r['left_sheet'], r['left_key'], r['right_sheet'], r['right_key'])
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    return out[:max_pairs]


def merge_sheets(sheets, left_sheet, left_key, right_sheet, right_key, how='left'):
    """Validated merge of two sheets. Returns (merged_df, join_report)."""
    import pandas as _pd
    if left_sheet not in sheets or right_sheet not in sheets:
        raise ValueError('Sheet not found.')
    left, right = sheets[left_sheet], sheets[right_sheet]
    if left_key not in left.columns:
        raise ValueError(f'Key "{left_key}" not found in "{left_sheet}".')
    if right_key not in right.columns:
        raise ValueError(f'Key "{right_key}" not found in "{right_sheet}".')
    if how not in ('left', 'inner'):
        raise ValueError('Join type must be "left" or "inner".')

    lkey = left[left_key].astype(str).str.strip().str.lower()
    rkey = right[right_key].astype(str).str.strip().str.lower()
    lblank = lkey.isin(['', 'nan', 'none', 'nat'])
    rblank = rkey.isin(['', 'nan', 'none', 'nat'])
    lkey = lkey.where(~lblank)
    rkey = rkey.where(~rblank)
    rset = set(rkey.dropna())
    matched_mask = lkey.isin(rset)
    dup_l = int(len(lkey.dropna()) - lkey.dropna().nunique())
    dup_r = int(len(rkey.dropna()) - rkey.dropna().nunique())

    L = left.copy()
    L['__join_key__'] = lkey
    R = right.copy()
    R['__join_key__'] = rkey
    merged = _pd.merge(L, R, on='__join_key__', how=how, suffixes=('', ' (right)'))
    merged = merged.drop(columns=['__join_key__'])

    join_report = {
        'left_sheet': left_sheet, 'left_key': left_key,
        'right_sheet': right_sheet, 'right_key': right_key,
        'join_type': how,
        'left_rows': int(len(left)), 'right_rows': int(len(right)),
        'matched_left_rows': int(matched_mask.sum()),
        'unmatched_left_rows': int((~matched_mask).sum()),
        'duplicate_left_keys': dup_l, 'duplicate_right_keys': dup_r,
        'result_rows': int(len(merged)),
        'result_columns': int(len(merged.columns)),
    }
    if dup_r:
        join_report['note'] = (f'"{right_key}" has {dup_r} duplicate value(s) in "{right_sheet}" - '
                               'the merge will multiply matching left rows (one-to-many). Verify before official use.')
    return merged, join_report


# ================================================================== #
# DATA VALIDATION - Excel entry rules generated from the actual data
# ================================================================== #
def build_validation_rules(cleaned, col_meta):
    """Generate Excel-style data validation rules from observed data. Rules only - no data changed."""
    rules = []
    for c in cleaned.columns:
        m = col_meta[c]
        role = m['role']
        if m.get('is_datetime'):
            s = cleaned[c].dropna()
            if len(s):
                rules.append({'column': c, 'type': 'date',
                              'min': str(s.min().date()), 'max': str(s.max().date()),
                              'rule': f'Date between {s.min().date()} and {s.max().date()}',
                              'purpose': 'Rejects mistyped / out-of-range dates at entry'})
        elif m.get('is_numeric'):
            s = cleaned[c].dropna()
            if not len(s):
                continue
            lo, hi = float(s.min()), float(s.max())
            if role == 'percent':
                rules.append({'column': c, 'type': 'decimal', 'min': 0, 'max': 100,
                              'rule': 'Decimal between 0 and 100 (percentage)',
                              'purpose': 'Blocks out-of-range percentages (e.g. 130%)'})
            elif role in ('amount', 'budget', 'expenditure'):
                pad = max(abs(lo), abs(hi), 1) * 0.10
                rules.append({'column': c, 'type': 'decimal',
                              'min': _j(lo - pad), 'max': _j(hi + pad),
                              'rule': f'Decimal between {_j(lo - pad):,} and {_j(hi + pad):,} (observed range ±10%)',
                              'purpose': 'Catches mistyped amounts (extra/missing zeros, wrong decimals)'})
            else:
                rules.append({'column': c, 'type': 'decimal', 'min': _j(lo), 'max': _j(hi),
                              'rule': f'Decimal between {_j(lo):,} and {_j(hi):,} (observed range)',
                              'purpose': 'Entry must stay within the observed data range'})
        elif cleaned[c].dtype == object:
            s = cleaned[c].dropna()
            if not len(s):
                continue
            uniq = s.astype(str).str.strip()
            uniq = uniq[uniq != '']
            nuniq = uniq.nunique()
            if role in CATEGORICAL_ROLES and 1 < nuniq <= 50:
                rules.append({'column': c, 'type': 'list',
                              'values': sorted(uniq.unique().tolist()),
                              'rule': f'Dropdown list - {nuniq} allowed value(s)',
                              'purpose': 'Prevents spelling variants like "finance" / "FINANCE" / "Finance "'})
            elif role == 'id':
                lens = uniq.str.len()
                if len(lens):
                    rules.append({'column': c, 'type': 'textLength',
                                  'min': int(lens.min()), 'max': int(lens.max()),
                                  'rule': f'Text length between {int(lens.min())} and {int(lens.max())} characters',
                                  'purpose': 'Catches truncated / malformed IDs'})
    return rules


# ================================================================== #
# POWER QUERY - M recipe generated from the actual cleaning steps
# ================================================================== #
def _m_str(v):
    return '"' + str(v).replace('"', '""') + '"'


def generate_powerquery(cleaned, col_meta, coerced_numeric_cols, options, ctx):
    """Generate a Power Query (M) script + Excel UI steps matching the agent's cleaning."""
    file_name = ctx.get('file_name', 'data.xlsx')
    sheet = ctx.get('sheet_name', 'Sheet1')
    is_csv = file_name.lower().endswith('.csv')
    has_dates = any(col_meta[c].get('is_datetime') for c in cleaned.columns)
    coerced = set(coerced_numeric_cols or [])
    std_case = options.get('standardize_case', True)
    fill_missing = options.get('fill_missing', True)

    type_pairs, trim_pairs, case_pairs, fill_pairs = [], [], [], []
    for c in cleaned.columns:
        m = col_meta[c]
        if m.get('is_datetime'):
            type_pairs.append((c, 'type date'))
        elif m.get('is_numeric'):
            allint = True
            s = cleaned[c].dropna()
            for v in (s.head(50) if len(s) else []):
                try:
                    if abs(float(v) - round(float(v))) > 1e-9:
                        allint = False
                        break
                except (TypeError, ValueError):
                    allint = False
                    break
            type_pairs.append((c, 'Int64.Type' if allint else 'type number'))
        else:
            type_pairs.append((c, 'type text'))
            if c in coerced:
                continue
            trim_pairs.append((c, 'each Text.Clean(Text.Trim(_)), type text'))
            if std_case and m['role'] in CATEGORICAL_ROLES:
                case_pairs.append((c, 'each Text.Proper(Text.Trim(_)), type text'))
            if fill_missing and m['role'] in CATEGORICAL_ROLES:
                fill_pairs.append((c, 'each if _ = null or Text.Trim(_) = "" then "MISSING" else _'))

    L = ['// =========================================================',
         '// GovData Analytics - Power Query script (generated)',
         f'// Source: {file_name} | Sheet: {sheet}',
         '// Adjust the file path below to match your computer.',
         '// =========================================================',
         'let']
    if is_csv:
        L.append(f'    Source = Csv.Document(File.Contents("C:\\Data\\{_m_str(file_name)[1:-1]}"), [Delimiter=",", Encoding=65001]),')
        L.append('    PromotedHeaders = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),')
        prev = 'PromotedHeaders'
    else:
        L.append(f'    Source = Excel.Workbook(File.Contents("C:\\Data\\{_m_str(file_name)[1:-1]}"), null, true),')
        L.append(f'    SheetData = Source{{[Item={_m_str(sheet)},Kind="Sheet"]}}[Data],')
        L.append('    PromotedHeaders = Table.PromoteHeaders(SheetData, [PromoteAllScalars=true]),')
        prev = 'PromotedHeaders'
    if type_pairs:
        pairs = ', '.join(f'{{{_m_str(c)}, {t}}}' for c, t in type_pairs)
        locale = ', "en-GB"' if has_dates else ''  # day-first dates
        L.append(f'    ChangedType = Table.TransformColumnTypes({prev}, {{{pairs}}}{locale}),')
        prev = 'ChangedType'
    for idx, col in enumerate(coerced, 1):
        L.append('    // "{}" contained text like "Rs. 45,000" - keep digits/dot/minus only'.format(col))
        step = 'CleanAmount{}'.format(idx)
        line = ('    ' + step + ' = Table.TransformColumns(' + prev + ', '
                '{{' + _m_str(col)
                + ', each try Number.From(Text.Select(_, {"0".."9", ".", "-"})) otherwise null, type number}}'
                + '),')
        L.append(line)
        prev = step
    if trim_pairs:
        pairs = ', '.join('{{{}, {}}}'.format(_m_str(c), f) for c, f in trim_pairs)
        L.append('    TrimmedText = Table.TransformColumns({}, {{{}}}),'.format(prev, pairs))
        prev = 'TrimmedText'
    if case_pairs:
        pairs = ', '.join('{{{}, {}}}'.format(_m_str(c), f) for c, f in case_pairs)
        L.append('    // Standardise department / district / status spelling to Title Case')
        L.append('    ProperCase = Table.TransformColumns({}, {{{}}}),'.format(prev, pairs))
        prev = 'ProperCase'
    if fill_pairs:
        pairs = ', '.join('{{{}, {}}}'.format(_m_str(c), f) for c, f in fill_pairs)
        L.append('    // Mark blank categories clearly instead of inventing values')
        L.append('    FilledMissing = Table.TransformColumns({}, {{{}}}),'.format(prev, pairs))
        prev = 'FilledMissing'
    L.append('    // Remove exact duplicate rows (delete this line if you must keep every row)')
    L.append(f'    RemovedDuplicates = Table.Distinct({prev})')
    L.append('in')
    L.append('    RemovedDuplicates')
    m_code = '\n'.join(L)

    ui_steps = [
        f'Excel kholein → Data → Get Data → {"From Text/CSV" if is_csv else "From File → From Workbook"} → apni file select karein.',
        f'Navigator me "{sheet}" select karein → "Transform Data" (Power Query Editor khulega).',
        'Home → Remove Rows → Remove Blank Rows (agar blank rows hain).',
    ]
    if has_dates:
        ui_steps.append('Date column select karein → data type icon → Using Locale → Date + English (United Kingdom) — day-first (dd/mm/yyyy) ke liye.')
    if coerced:
        ui_steps.append('Text amount columns ("Rs. 45,000" wale) select karke Transform tab se data type ko Decimal Number karein.')
    if trim_pairs:
        ui_steps.append('Text columns select → Transform → Format → Trim (extra spaces hat jayenge).')
    if std_case:
        ui_steps.append('Department/District/Status columns select → Transform → Format → Capitalize Each Word.')
    ui_steps.append('Home → Remove Duplicates (sirf tab jab exact duplicate rows confirm ho chuke hain).')
    ui_steps.append('Home → Close & Load To… → "PivotTable Report" — ab pivot khud banayein, ya "Table" par load karein.')
    ui_steps.append('Refresh: Data → Refresh All (naya data aane par sirf refresh karna hota hai — formulas dobara banane ki zaroorat nahi).')
    return {'m_code': m_code, 'ui_steps': ui_steps, 'file_name': file_name, 'sheet': sheet}


# ================================================================== #
# MANAGEMENT SUMMARY - auto executive summary (FACT / OBSERVATION style)
# ================================================================== #
def build_management_summary(kpis, pivots, budget, issues, trend, top50):
    """Auto-generate a concise management summary. Only states what the data shows."""
    lines = []
    total = kpis.get('total_records')
    if total is not None:
        lines.append(f"FACT: Dataset contains {total:,} records across {kpis.get('columns', '?')} fields"
                     + (f" (sheet \"{kpis.get('sheet_name', '')}\")." if kpis.get('sheet_name') else '.'))
    if kpis.get('date_range'):
        lines.append(f"FACT: Reference period {kpis['date_range'][0]} to {kpis['date_range'][1]}.")
    if budget:
        lines.append(f"FACT: Total budget {_fmt_big(budget['total_budget'])}; expenditure {_fmt_big(budget['total_expenditure'])}.")
        if budget.get('utilization') is not None:
            lines.append(f"CALCULATION: Utilisation {budget['utilization']}% "
                         f"(expenditure / budget). Remaining: {_fmt_big(budget.get('remaining'))}.")
        if budget.get('above_budget_count'):
            lines.append(f"POTENTIAL ISSUE: {budget['above_budget_count']} record(s) where expenditure exceeds the budget allocation - requires review.")
    if kpis.get('main_metric') and kpis.get('main_metric_total') is not None:
        lines.append(f"CALCULATION: Total {kpis['main_metric']} = {_fmt_big(kpis['main_metric_total'])}.")
    if pivots:
        p0 = pivots[0]
        if p0['rows']:
            top = p0['rows'][0]
            share = f" ({top['pct']}% of total)" if top.get('pct') is not None and p0.get('has_measure') else f" ({round(top['count'] / max(total or 1, 1) * 100, 1)}% of records)"
            lines.append(f"OBSERVATION: Largest {p0['dimension']} is \"{top['label']}\"{share}.")
    if trend and len(trend.get('labels', [])) >= 2:
        vals = [v for v in trend['values'] if v is not None]
        if len(vals) >= 2 and vals[0]:
            last, first = vals[-1], vals[0]
            chg = (last - first) / abs(first) * 100 if first else None
            if chg is not None:
                direction = 'increased' if chg > 0 else ('decreased' if chg < 0 else 'stayed flat')
                lines.append(f"OBSERVATION: Monthly {trend['measure'] or 'record count'} {direction} from {trend['labels'][0]} ({_fmt_big(first)}) "
                             f"to {trend['labels'][-1]} ({_fmt_big(last)}) - change {chg:+.1f}% over the period.")
    if kpis.get('duplicate_rows'):
        lines.append(f"POTENTIAL ISSUE: {kpis['duplicate_rows']:,} exact duplicate record(s) - kept for audit, not deleted.")
    if kpis.get('missing_cells'):
        lines.append(f"DATA LIMITATION: {kpis['missing_cells']:,} missing cells ({kpis.get('missing_pct')}% of all cells). No values were invented.")
    if kpis.get('outliers'):
        lines.append(f"POTENTIAL ISSUE: {kpis['outliers']} statistical anomal{('y' if kpis['outliers'] == 1 else 'ies')} flagged (IQR) - requires review, not proof of error.")
    lines.append("NOTE: Every figure above is computed directly from the uploaded file; nothing is estimated.")
    return lines


def _fmt_big(v):
    if v is None:
        return 'not available'
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    av = abs(v)
    if av >= 1e9:
        return f"{v / 1e9:,.2f}B"
    if av >= 1e6:
        return f"{v / 1e6:,.2f}M"
    if av >= 1e3:
        return f"{v / 1e3:,.1f}K"
    return f"{v:,.0f}"


# ================================================================== #
# EXCEL FORMULA RECIPES - VLOOKUP / XLOOKUP / SUMIFS etc. from the
# user's ACTUAL workbook structure (real sheet names, ranges, columns)
# ================================================================== #
def _xl_letter(idx):
    """1-based column index -> Excel letter."""
    letter = ''
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        letter = chr(65 + rem) + letter
    return letter


def build_formula_recipes(raw, col_meta, ctx, main_metric, dim_cols, id_dup_info):
    """Generate copy-paste Excel formulas using the user's real sheet/column names."""
    recipes = []
    cols = list(raw.columns)
    n = len(raw)
    sheet = ctx.get('sheet_name', 'Sheet1')
    q = '' if str(sheet).replace(' ', '').isalnum() else "'"
    sheet_ref = f'{q}{sheet}{q}'

    def L(col):
        return _xl_letter(cols.index(col) + 1)

    key_col = next((c for c in cols if col_meta[c]['role'] == 'id'), None)
    if key_col is None and dim_cols:
        key_col = dim_cols[0]
    if key_col is None:
        key_col = cols[0]
    key_letter, key_last = L(key_col), n + 1

    # ---------- cross-sheet lookup candidates (any shared column, ID/dimension preferred) ----------
    wb_sheets = ctx.get('workbook_sheets') or {}
    col_priority = ([key_col] if key_col else []) + [c for c in dim_cols if c != key_col] + \
                   [c for c in cols if c not in ([key_col] if key_col else []) and c not in dim_cols]
    candidates = []
    for cand_col in col_priority:
        for sname, sinfo in wb_sheets.items():
            if sname == sheet:
                continue
            scol = sinfo.get('columns', [])
            shared = [c for c in scol if str(c).strip().lower() == str(cand_col).strip().lower()]
            if not shared:
                continue
            returns = [c for c in scol if c != shared[0]]
            if not returns:
                continue
            candidates.append((cand_col, sname, shared[0], int(sinfo.get('rows', 0)), scol, returns[0]))
            break   # one lookup sheet per key column
        if len(candidates) >= 2:
            break
    for cand_col, ls_name, lkey, lrows, lcols, ret0 in candidates[:2]:
        lq = '' if str(ls_name).replace(' ', '').isalnum() else "'"
        lref = f'{lq}{ls_name}{lq}'
        cand_letter = _xl_letter(cols.index(cand_col) + 1)
        lk_letter = _xl_letter(lcols.index(lkey) + 1)
        r_letter = _xl_letter(lcols.index(ret0) + 1)
        vidx = lcols.index(ret0) + 1
        rng_k = f'{lref}!${lk_letter}$2:${lk_letter}${lrows + 1}'
        rng_r = f'{lref}!${r_letter}$2:${r_letter}${lrows + 1}'
        rng_all = f'{lref}!$A$2:${_xl_letter(len(lcols))}${lrows + 1}'
        tag = f'"{ls_name}" se "{ret0}" laao (key: "{cand_col}" = "{lkey}")'
        recipes += [
            {'section': f'XLOOKUP - {ls_name} (Excel 365 / 2021+)',
             'formula': f'=XLOOKUP({cand_letter}2, {rng_k}, {rng_r}, "Not Found")',
             'purpose': tag,
             'explain': (f'{cand_letter}2 = jis ki value dhoondni hai ("{cand_col}"); {rng_k} = lookup sheet ka key column; '
                         f'{rng_r} = wahi column jo wapas chahiye ("{ret0}"); "Not Found" = jab match na ho. Exact match by default.')},
            {'section': f'VLOOKUP - {ls_name} (sab versions)',
             'formula': f'=VLOOKUP({cand_letter}2, {rng_all}, {vidx}, FALSE)',
             'purpose': tag + ' - purane Excel ke liye',
             'explain': (f'{rng_all} = lookup table (key column FIRST hona chahiye - VLOOKUP ki limitation); '
                         f'{vidx} = "{ret0}" ka column number table ke andar; FALSE = exact match (hamesha FALSE rakhein).')},
            {'section': f'INDEX + MATCH - {ls_name} (purane Excel, flexible)',
             'formula': f'=INDEX({rng_r}, MATCH({cand_letter}2, {rng_k}, 0))',
             'purpose': 'VLOOKUP ka behtar alternative - left-right dono taraf dekh sakta hai',
             'explain': 'MATCH key ki position dhoondta hai (0 = exact); INDEX us position se value laata hai.'},
            {'section': f'IFERROR + VLOOKUP - {ls_name} (#N/A handle)',
             'formula': f'=IFERROR(VLOOKUP({cand_letter}2, {rng_all}, {vidx}, FALSE), "Not Found")',
             'purpose': 'Not-found par #N/A ki jagah saaf-saaf "Not Found" likhega',
             'explain': 'IFERROR andar wale formula ke error ko doosri value se replace karta hai.'},
        ]
    if not candidates and main_metric and key_col:
        # within-sheet example: bring main metric by key
        m_letter = L(main_metric)
        recipes.append({
            'section': 'XLOOKUP - same sheet example (Excel 365 / 2021+)',
            'formula': f'=XLOOKUP({key_letter}2, {sheet_ref}!${key_letter}$2:${key_letter}${key_last}, '
                       f'{sheet_ref}!${m_letter}$2:${m_letter}${key_last}, "Not Found")',
            'purpose': f'Key "{key_col}" se "{main_metric}" dhoondne ka example (isi sheet ke andar)',
            'explain': 'Multi-sheet workbook upload karein to cross-sheet formulas khud ban jayengi.'})

    # ---------- conditional aggregations ----------
    dim = dim_cols[0] if dim_cols else None
    if main_metric and dim:
        d_letter, m_letter = L(dim), L(main_metric)
        sample = raw[dim].dropna()
        sample_val = str(sample.iloc[0]) if len(sample) else 'X'
        rng_d = f'{sheet_ref}!${d_letter}:${d_letter}'
        rng_m = f'{sheet_ref}!${m_letter}:${m_letter}'
        recipes += [
            {'section': 'SUMIFS - department-wise total',
             'formula': f'=SUMIFS({rng_m}, {rng_d}, "{sample_val}")',
             'purpose': f'Sirf "{sample_val}" ke records ka {main_metric} total',
             'explain': f'{rng_m} = jis column ka sum karna hai; {rng_d} + "{sample_val}" = condition (kis department ka).'},
            {'section': 'COUNTIFS - record count with condition',
             'formula': f'=COUNTIFS({rng_d}, "{sample_val}")',
             'purpose': f'"{sample_val}" ke kitne records hain',
             'explain': 'Sirf ginta hai - koi sum nahi.'},
            {'section': 'AVERAGEIFS - department-wise average',
             'formula': f'=AVERAGEIFS({rng_m}, {rng_d}, "{sample_val}")',
             'purpose': f'"{sample_val}" ka average {main_metric}',
             'explain': 'SUMIFS jaisa, lekin total ki jagah average.'},
            {'section': 'MAXIFS / MINIFS (Excel 2019+)',
             'formula': f'=MAXIFS({rng_m}, {rng_d}, "{sample_val}")',
             'purpose': f'"{sample_val}" ka sab se bara {main_metric}',
             'explain': 'MINIFS iska ulta hai. Purane Excel ke liye array formula chahiye hota hai.'},
        ]
    # ---------- ranking ----------
    if main_metric:
        m_letter = L(main_metric)
        recipes.append({
            'section': 'RANK.EQ - ranking',
            'formula': f'=RANK.EQ({m_letter}2, {sheet_ref}!${m_letter}$2:${m_letter}${key_last}, 0)',
            'purpose': f'{main_metric} ki ranking (highest = 1)',
            'explain': '0 (ya blank) = descending (sab se bara #1); 1 = ascending. Ties par same rank milta hai (competition ranking).'})
    return recipes


# ------------------------------------------------------------------ #
# HIGHLIGHT & SUMMARIZE DUPLICATES (VBA feature ka Python port)
# Key: Name + Age + Department + Salary — Trim ke baad EXACT match,
# 1 se zyada baar = duplicate. Name akela kaafi nahi. Raw values use
# hote hain (cleaning se pehle), taake VBA ke ActiveSheet jaisa rahe.
# ------------------------------------------------------------------ #
DUP_COMBO_REQUIRED = [('name', 'Name'), ('age', 'Age'),
                      ('department', 'Department'), ('salary', 'Salary')]


def _vba_cstr(v):
    """VBA ka Trim(CStr(value)) behaviour — duplicate key ke liye exact port.

    - None/NaN -> ""          (VBA: Empty -> "")
    - 30.0 -> "30"            (VBA CStr(30#) -> "30")
    - strings par sirf spaces trim hoti hain (VBA Trim sirf 0x20 katta hai)
    """
    if v is None:
        return ''
    if isinstance(v, float):
        if v != v:  # NaN
            return ''
        return str(int(v)) if v.is_integer() else ('%g' % v)
    if isinstance(v, bool):
        return 'True' if v else 'False'
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    if isinstance(v, (float, np.floating)):
        f = float(v)
        return str(int(f)) if f.is_integer() else ('%g' % f)
    if isinstance(v, str):
        return v.strip(' ')
    if 'Timestamp' in type(v).__name__ or isinstance(v, (np.datetime64,)):
        return str(v)
    return str(v).strip(' ')


def _dup_salary_num(v):
    """Salary ko number me parse karo (VBA IsNumeric/CDbl jaisa). None = numeric nahi."""
    if isinstance(v, (bool, np.bool_)):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        f = float(v)
        return None if f != f else f
    if isinstance(v, str) and v.strip():
        try:
            return float(v.replace(',', ''))
        except ValueError:
            return None
    return None


def _dup_val(v):
    """Cell value ko display/JSON-safe banao (types preserve, NaN->None)."""
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:
            return None
        return int(v) if v.is_integer() else v
    if isinstance(v, (bool, np.bool_)):
        return 'YES' if v else 'NO'
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (np.floating,)):
        f = float(v)
        return int(f) if f.is_integer() else f
    if isinstance(v, str):
        return v
    return str(v)


def analyze_duplicate_combos(raw):
    """Duplicate combinations detect karo (VBA HighlightAndSummarizeDuplicates wahi logic).

    Returns dict: ok, columns_used, combos[] (name/age/department/salary/count/
    total_salary/row_numbers), dup_row_numbers (1-based data rows), validation{},
    message. Total Salary = Salary x Duplicate Count (spec rule).
    """
    out = {'feature': 'highlight_duplicates',
           'key': ['Name', 'Age', 'Department', 'Salary'],
           'match': 'exact after Trim (VBA Trim/CStr jaisa) — case-sensitive, Name akela nahi'}
    if raw is None or not len(list(raw.columns)):
        out.update(ok=False, error='Is sheet me koi columns nahi mile.')
        return out

    # required headers (trimmed, case-insensitive) — VBA version B/C/D/E positions use karta hai
    hdr = {}
    for c in raw.columns:
        h = str(c).strip().lower()
        if h and h not in hdr:
            hdr[h] = c
    cols, missing = {}, []
    for k, label in DUP_COMBO_REQUIRED:
        if k in hdr:
            cols[k] = hdr[k]
        else:
            missing.append(label)
    if missing:
        out.update(ok=False,
                   error=('Required columns missing: ' + ', '.join(missing) +
                          '. Ye feature headers "Name", "Age", "Department" aur "Salary" maangta hai '
                          '(VBA version me Column B=Name, C=Age, D=Department, E=Salary).'),
                   missing=missing)
        return out

    name_s = raw[cols['name']].tolist()
    age_s = raw[cols['age']].tolist()
    dep_s = raw[cols['department']].tolist()
    sal_s = raw[cols['salary']].tolist()

    # ---- FIRST PASS: har combination kitni baar aaya (VBA dict count) ----
    keys = ['|'.join((_vba_cstr(a), _vba_cstr(b), _vba_cstr(c), _vba_cstr(d)))
            for a, b, c, d in zip(name_s, age_s, dep_s, sal_s)]
    counts = {}
    for k in keys:
        counts[k] = counts.get(k, 0) + 1
    dup_keys = {k for k, n in counts.items() if n > 1}

    # ---- SECOND PASS: groups banao (first-occurrence order, VBA jaisa) ----
    combos, seen = [], {}
    dup_rows = []
    non_numeric = 0
    for i, k in enumerate(keys, 1):
        if k not in dup_keys:
            continue
        dup_rows.append(i)
        if k not in seen:
            seen[k] = len(combos)
            sal_raw = sal_s[i - 1]
            sal_num = _dup_salary_num(sal_raw)
            if sal_num is None:
                non_numeric += counts[k]
            combos.append({'name': _dup_val(name_s[i - 1]), 'age': _dup_val(age_s[i - 1]),
                           'department': _dup_val(dep_s[i - 1]), 'salary': _dup_val(sal_raw),
                           'salary_num': sal_num, 'count': counts[k], 'row_numbers': [i]})
        else:
            combos[seen[k]]['row_numbers'].append(i)

    # Total Salary = Salary x Duplicate Count (spec rule — unrelated salaries add NAHI hote)
    for c in combos:
        c['total_salary'] = c['salary_num'] * c['count'] if c['salary_num'] is not None else None
        c.pop('salary_num', None)

    # sab se zyada repeat pehle (display ke liye)
    combos.sort(key=lambda c: (-c['count'], str(c['name'])))

    # ---- VALIDATION (independent recount — fabricated numbers kabhi nahi) ----
    v = {
        'dup_rows_found': len(dup_rows),
        'unique_combos': len(combos),
        'summary_rows': len(combos),
        'total_salary_check': True,
        'highlight_rows_all_in_groups': all(counts[keys[i - 1]] > 1 for i in dup_rows),
    }
    v['recount_independent'] = sum(1 for k in keys if counts[k] > 1)
    for c in combos:  # Total Salary ko dobara, alag tareeqe se check karo
        if c['total_salary'] is None:
            continue
        ssum = 0.0
        for i in c['row_numbers']:
            sv = _dup_salary_num(sal_s[i - 1])
            if sv is not None:
                ssum += sv
        if abs(ssum - c['total_salary']) > 0.005:
            v['total_salary_check'] = False
    v['no_nondup_highlighted'] = (v['highlight_rows_all_in_groups']
                                  and v['recount_independent'] == len(dup_rows))
    v['source_values_untouched'] = True  # report likhne ke BAAD app.py asal values se verify karta hai

    if dup_rows:
        message = (f'Duplicate analysis completed. {len(dup_rows)} duplicate rows found. '
                   f'{len(combos)} unique duplicate entries created. '
                   "A summary has been created in the 'Duplicates' sheet.")
    else:
        message = 'No duplicate records were found.'

    out.update(ok=True, columns_used={k: str(c) for k, c in cols.items()},
               total_rows=len(raw), dup_row_count=len(dup_rows), unique_combos=len(combos),
               dup_row_numbers=dup_rows, combos=combos, non_numeric_salary=non_numeric,
               validation=v, message=message)
    return out
